"""Phase 4 tests: opera_cloud plugin backend (models, routes, export, expiry).

Uses a throwaway Postgres schema; the real plugins/ dir is loaded
(hello + opera_cloud). Follows the test_plugins.py pattern.
"""
import os
import sys
import uuid
from datetime import date, timedelta
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, text
from sqlalchemy.orm import sessionmaker

os.environ.setdefault(
    "DATABASE_URL",
    "postgresql+psycopg2://registry:registry_dev_only@localhost:5433/account_registry",
)

REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))
REPO_PLUGINS = REPO_ROOT / "plugins"
if not REPO_PLUGINS.is_dir():
    for cand in (REPO_ROOT / "app" / "plugins", Path("/app/plugins")):
        if cand.is_dir():
            REPO_PLUGINS = cand
            break

from app.core import auth, models  # noqa: E402
from app.core import plugin_loader  # noqa: E402
from app.core.db import Base, get_db  # noqa: E402

import plugins.opera_cloud.models  # noqa: E402,F401 — registers OperaAccount
import plugins.opera_cloud.services as opera_services  # noqa: E402

TEST_URL = os.environ["DATABASE_URL"]

OPERA_PERMS = {
    "opera.account.view", "opera.account.create", "opera.account.edit",
    "opera.account.delete", "opera.account.export",
}


@pytest.fixture(scope="module")
def opera_app():
    schema = f"o_{uuid.uuid4().hex[:10]}"
    admin = create_engine(TEST_URL)
    with admin.connect() as c:
        c.execute(text(f"CREATE SCHEMA {schema}"))
        c.commit()
    url = f"{TEST_URL}?options=-csearch_path%3D{schema}"
    test_engine = create_engine(url)
    maker = sessionmaker(bind=test_engine)
    Base.metadata.create_all(test_engine)

    db = maker()
    db.add(models.Role(name="admin", description="Full access"))
    db.commit()
    admin_role = db.query(models.Role).filter_by(name="admin").one()
    db.add(models.User(
        username="admin", email="a@x.io",
        password_hash=auth.hash_password("Admin123!"), roles=[admin_role],
    ))
    db.commit()
    db.close()

    app = FastAPI()
    from app.core.routes import router as core_api_router
    app.include_router(core_api_router)

    def override():
        db = maker()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override
    loaded = plugin_loader.load_plugins(
        app, session_factory=maker, db_engine=test_engine,
        plugins_dir=REPO_PLUGINS,
    )
    assert "opera_cloud" in loaded, f"opera_cloud not loaded: {plugin_loader.SKIPPED}"

    client = TestClient(app, raise_server_exceptions=False)
    yield client, maker, test_engine
    app.dependency_overrides.clear()
    plugin_loader.LOADED.clear()
    plugin_loader.SKIPPED.clear()
    plugin_loader._MOUNTED.clear()
    Base.metadata.drop_all(test_engine)
    with admin.connect() as c:
        c.execute(text(f"DROP SCHEMA {schema} CASCADE"))
        c.commit()


def _login(client, username="admin", password="Admin123!"):
    r = client.post("/api/v1/core/auth/login",
                    json={"username": username, "password": password})
    assert r.status_code == 200, r.text
    toks = r.json()
    return {"Authorization": f"Bearer {toks['access_token']}"}


def _payload(n=1, **kw):
    base = {
        "username": f"op{n}",
        "email": f"op{n}@x.io",
        "employee_name": f"Emp {n}",
        "employee_id": f"E{n:04d}",
        "operator_code": f"OP{n:04d}",
        "roles": ["cashier"],
        "privileges": ["checkin"],
    }
    base.update(kw)
    return base


def test_plugin_listed_menu_perms(opera_app):
    client, maker, _ = opera_app
    h = _login(client)
    r = client.get("/api/v1/core/plugins", headers=h)
    assert r.status_code == 200
    opera = [p for p in r.json() if p["name"] == "opera_cloud"]
    assert opera and opera[0]["version"] == "1.0.0"
    assert set(opera[0]["permissions"]) == OPERA_PERMS
    r = client.get("/api/v1/core/menu", headers=h)
    assert r.status_code == 200
    assert any(i["path"] == "/opera" for i in r.json())
    db = maker()
    codes = {p.code for p in db.query(models.Permission).all()}
    db.close()
    assert OPERA_PERMS <= codes


def test_migration_table_exists(opera_app):
    _, _, test_engine = opera_app
    with test_engine.connect() as c:
        names = {row[0] for row in c.execute(text(
            "SELECT tablename FROM pg_tables WHERE schemaname = current_schema()"))}
        ver = c.execute(text("SELECT version_num FROM alembic_version_opera_cloud")).fetchall()
    assert "opera_accounts" in names
    assert ver and ver[0][0] == "opera_0001"


def test_rbac(opera_app):
    client, maker, _ = opera_app
    assert client.get("/api/v1/plugins/opera_cloud/accounts").status_code == 401
    db = maker()
    db.add(models.User(username="noview", email="nv@x.io",
                       password_hash=auth.hash_password("Noview123!")))
    db.commit()
    db.close()
    r = client.post("/api/v1/core/auth/login",
                    json={"username": "noview", "password": "Noview123!"})
    vh = {"Authorization": f"Bearer {r.json()['access_token']}"}
    assert client.get("/api/v1/plugins/opera_cloud/accounts", headers=vh).status_code == 403
    h = _login(client)
    assert client.get("/api/v1/plugins/opera_cloud/accounts", headers=h).status_code == 200


def test_crud_search_filters_audit(opera_app):
    client, maker, _ = opera_app
    h = _login(client)
    # seed dept/position for FK checks
    d = client.post("/api/v1/core/departments", headers=h, json={"name": "FO"}).json()
    p = client.post("/api/v1/core/positions", headers=h,
                    json={"name": "Recep", "department_id": d["id"]}).json()
    r = client.post("/api/v1/plugins/opera_cloud/accounts", headers=h,
                    json=_payload(11, department_id=d["id"], position_id=p["id"]))
    assert r.status_code == 201, r.text
    acc = r.json()
    assert acc["department_name"] == "FO" and acc["position_name"] == "Recep"
    aid = acc["id"]
    client.post("/api/v1/plugins/opera_cloud/accounts", headers=h, json=_payload(12))
    # q search across all five fields
    for q in ("op11", "op11@x.io", "Emp 11", "E0011", "OP0011"):
        r = client.get("/api/v1/plugins/opera_cloud/accounts",
                       headers=h, params={"q": q})
        assert r.status_code == 200 and r.json()["total"] >= 1, q
    # filters
    r = client.get("/api/v1/plugins/opera_cloud/accounts", headers=h,
                   params={"department_id": d["id"]})
    assert r.status_code == 200 and all(
        a["department_id"] == d["id"] for a in r.json()["items"])
    r = client.get("/api/v1/plugins/opera_cloud/accounts", headers=h,
                   params={"account_status": "active"})
    assert r.status_code == 200 and r.json()["total"] >= 2
    # pagination shape
    r = client.get("/api/v1/plugins/opera_cloud/accounts", headers=h,
                   params={"page": 2, "page_size": 1})
    body = r.json()
    assert set(body) == {"items", "total", "page", "page_size"}
    # get one + update + audit trail
    assert client.get(f"/api/v1/plugins/opera_cloud/accounts/{aid}",
                      headers=h).status_code == 200
    r = client.put(f"/api/v1/plugins/opera_cloud/accounts/{aid}", headers=h,
                   json={"employee_name": "Emp Eleven", "account_status": "disabled"})
    assert r.status_code == 200 and r.json()["employee_name"] == "Emp Eleven"
    r = client.get("/api/v1/plugins/opera_cloud/accounts", headers=h,
                   params={"account_status": "disabled"})
    assert r.json()["total"] >= 1
    # delete
    assert client.delete(f"/api/v1/plugins/opera_cloud/accounts/{aid}",
                         headers=h).status_code == 204
    assert client.get(f"/api/v1/plugins/opera_cloud/accounts/{aid}",
                      headers=h).status_code == 404
    db = maker()
    rows = db.query(models.AuditLog).filter(
        models.AuditLog.resource == "opera_account").all()
    db.close()
    actions = {x.action for x in rows}
    assert {"create", "update", "delete"} <= actions


def test_validation_conflicts(opera_app):
    client, _, _ = opera_app
    h = _login(client)
    client.post("/api/v1/plugins/opera_cloud/accounts", headers=h, json=_payload(21))
    # duplicate username → 409
    r = client.post("/api/v1/plugins/opera_cloud/accounts", headers=h,
                    json=_payload(22, username="op21"))
    assert r.status_code == 409, r.text
    # duplicate operator_code → 409
    r = client.post("/api/v1/plugins/opera_cloud/accounts", headers=h,
                    json=_payload(23, operator_code="OP0021"))
    assert r.status_code == 409
    # bad email → 422
    r = client.post("/api/v1/plugins/opera_cloud/accounts", headers=h,
                    json=_payload(24, email="not-an-email"))
    assert r.status_code == 422
    # bad FK → 404
    r = client.post("/api/v1/plugins/opera_cloud/accounts", headers=h,
                    json=_payload(25, department_id=999999))
    assert r.status_code == 404
    # missing required → 422
    bad = _payload(26)
    del bad["operator_code"]
    assert client.post("/api/v1/plugins/opera_cloud/accounts",
                       headers=h, json=bad).status_code == 422


def test_export_xlsx(opera_app):
    client, _, _ = opera_app
    h = _login(client)
    client.post("/api/v1/plugins/opera_cloud/accounts", headers=h, json=_payload(31))
    r = client.get("/api/v1/plugins/opera_cloud/accounts/export", headers=h)
    assert r.status_code == 200, r.text
    assert "spreadsheetml.sheet" in r.headers["content-type"]
    assert r.content[:2] == b"PK"  # xlsx is a zip
    from openpyxl import load_workbook
    from io import BytesIO
    wb = load_workbook(BytesIO(r.content))
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert header == opera_services.EXPORT_COLUMNS
    assert ws.max_row >= 2
    # filter respected: export with q matching nothing has header only
    r = client.get("/api/v1/plugins/opera_cloud/accounts/export", headers=h,
                   params={"q": "zzz-no-such-user"})
    wb2 = load_workbook(BytesIO(r.content))
    assert wb2.active.max_row == 1


def test_expire_job(opera_app):
    client, maker, _ = opera_app
    h = _login(client)
    yesterday = (date.today() - timedelta(days=1)).isoformat()
    r = client.post("/api/v1/plugins/opera_cloud/accounts", headers=h,
                    json=_payload(41, expiration_date=yesterday))
    assert r.status_code == 201, r.text
    aid = r.json()["id"]
    db = maker()
    n = opera_services.expire_old_accounts(db)
    db.close()
    assert n >= 1
    r = client.get(f"/api/v1/plugins/opera_cloud/accounts/{aid}", headers=h)
    assert r.json()["account_status"] == "expired"
