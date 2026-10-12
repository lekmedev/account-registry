"""{{name}} plugin backend tests.

Plugin tests live here, next to the plugin — writing a plugin never means
editing backend/tests/. This is also the template new plugins copy.

Fixture and helpers come from core; see docs/PLUGIN-AUTHORING.md.
"""
import pytest
from sqlalchemy import text

from app.core import auth
from app.core import models as core_models

from tests.conftest import REPO_PLUGINS, app_factory, login  # noqa: F401

API = "/api/v1/plugins/{{name}}/records"

PERMS = {
    "{{PERM_PREFIX}}.record.view",
    "{{PERM_PREFIX}}.record.create",
    "{{PERM_PREFIX}}.record.edit",
    "{{PERM_PREFIX}}.record.delete",
    "{{PERM_PREFIX}}.record.export",
}


@pytest.fixture(scope="module")
def plug(app_factory):
    return app_factory("{{name}}")


def _payload(n=1, **kw):
    base = {
        "code": f"rec{n}",
        "name": f"Record {n}",
        "email": f"rec{n}@x.io",
        "status": "active",
        "tags": ["alpha"],
    }
    base.update(kw)
    return base


def test_plugin_listed_menu_perms(plug):
    client, maker, _ = plug
    h = login(client)
    r = client.get("/api/v1/core/plugins", headers=h)
    assert r.status_code == 200
    mine = [p for p in r.json() if p["name"] == "{{name}}"]
    assert mine, f"{{name}} not listed: {r.text}"
    assert set(mine[0]["permissions"]) == PERMS
    r = client.get("/api/v1/core/menu", headers=h)
    assert r.status_code == 200
    assert any(i["path"] == "{{ROUTE_PATH}}" for i in r.json())
    db = maker()
    codes = {p.code for p in db.query(core_models.Permission).all()}
    db.close()
    assert PERMS <= codes


def test_migration_ran(plug):
    _, _, engine = plug
    with engine.connect() as c:
        names = {row[0] for row in c.execute(text(
            "SELECT tablename FROM pg_tables WHERE schemaname = current_schema()"))}
        ver = c.execute(
            text("SELECT version_num FROM alembic_version_{{name}}")).fetchall()
    assert "{{TABLE}}" in names
    assert ver and ver[0][0] == "{{name}}_0001"


def test_rbac(plug):
    client, maker, _ = plug
    assert client.get(API).status_code == 401
    db = maker()
    db.add(core_models.User(username="noview", email="nv@x.io",
                            password_hash=auth.hash_password("Noview123!")))
    db.commit()
    db.close()
    vh = login(client, "noview", "Noview123!")
    assert client.get(API, headers=vh).status_code == 403
    assert client.get(API, headers=login(client)).status_code == 200


def test_crud_search_filter_audit(plug):
    client, maker, _ = plug
    h = login(client)

    r = client.post(API, headers=h, json=_payload(11, status="active"))
    assert r.status_code == 201, r.text
    rec_id = r.json()["id"]

    client.post(API, headers=h, json=_payload(12))

    # q search hits every text column the service declares searchable.
    # tags is JSONB — searching it needs its own operator, so it is out of
    # scope here; add a case when you implement it.
    for q in ("rec11", "rec11@x.io", "Record 11"):
        r = client.get(API, headers=h, params={"q": q})
        assert r.status_code == 200 and r.json()["total"] >= 1, q

    # filter
    r = client.get(API, headers=h, params={"status": "active"})
    assert r.status_code == 200 and all(i["status"] == "active" for i in r.json()["items"])

    # paging shape
    r = client.get(API, headers=h, params={"page": 2, "page_size": 1})
    assert set(r.json()) == {"items", "total", "page", "page_size"}

    # get one + partial update (PUT is partial, like core)
    assert client.get(f"{API}/{rec_id}", headers=h).status_code == 200
    r = client.put(f"{API}/{rec_id}", headers=h, json={"name": "Renamed"})
    assert r.status_code == 200 and r.json()["name"] == "Renamed"
    assert r.json()["code"] == "rec11"  # untouched fields survive

    # delete → 204, then 404
    assert client.delete(f"{API}/{rec_id}", headers=h).status_code == 204
    assert client.get(f"{API}/{rec_id}", headers=h).status_code == 404

    db = maker()
    rows = db.query(core_models.AuditLog).filter(
        core_models.AuditLog.resource == "{{name}}_record").all()
    db.close()
    assert {"create", "update", "delete"} <= {x.action for x in rows}


def test_validation_and_conflicts(plug):
    client, _, _ = plug
    h = login(client)
    client.post(API, headers=h, json=_payload(21))

    # duplicate unique code → 409
    r = client.post(API, headers=h, json=_payload(22, code="rec21"))
    assert r.status_code == 409, r.text
    # bad email → 422
    r = client.post(API, headers=h, json=_payload(23, email="not-an-email"))
    assert r.status_code == 422
    # missing required field → 422
    bad = _payload(24)
    del bad["code"]
    assert client.post(API, headers=h, json=bad).status_code == 422
    # unknown id → 404
    assert client.get(f"{API}/999999", headers=h).status_code == 404


def test_export_xlsx(plug):
    client, _, _ = plug
    h = login(client)
    client.post(API, headers=h, json=_payload(31))
    r = client.get(f"{API}/export", headers=h)
    assert r.status_code == 200, r.text
    assert "spreadsheetml.sheet" in r.headers["content-type"]
    assert r.content[:2] == b"PK"  # xlsx is a zip
    from openpyxl import load_workbook
    import io
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.cell(row=1, column=1).value == "id"
    assert ws.max_row >= 2


def test_task_registered():
    """register_tasks() must accept the scheduler core hands it."""
    import importlib.util
    import sys

    mod_path = REPO_PLUGINS / "{{name}}" / "tasks.py"
    spec = importlib.util.spec_from_file_location("{{name}}_tasks_probe", mod_path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = mod
    spec.loader.exec_module(mod)

    class _FakeScheduler:
        def __init__(self):
            self.jobs = []

        def add_job(self, fn, trigger, **kw):
            self.jobs.append((fn, trigger, kw))

    sched = _FakeScheduler()
    mod.register_tasks(sched)
    assert len(sched.jobs) == 1
    assert sched.jobs[0][1] == "cron"