"""opera_cloud backend tests (models, routes, export, expiry).

Lives next to the plugin — see docs/PLUGIN-AUTHORING.md. The fixture that
creates the throwaway schema and boots the app comes from core conftest.
"""
import io
from datetime import date, timedelta

import pytest
from sqlalchemy import text

from app.core import auth
from app.core import models as core_models

from tests.conftest import app_factory, login  # noqa: F401

import plugins.opera_cloud.models  # noqa: F401 — registers OperaAccount
import plugins.opera_cloud.services as opera_services

API = "/api/v1/plugins/opera_cloud/accounts"

OPERA_PERMS = {
    "opera.account.view", "opera.account.create", "opera.account.edit",
    "opera.account.delete", "opera.account.export",
}


@pytest.fixture(scope="module")
def opera_app(app_factory):
    return app_factory("opera_cloud")


def _payload(n=1, **kw):
    base = {
        "username": f"op{n}",
        "email": f"op{n}@x.io",
        "employee_name": f"Emp {n}",
        "employee_id": f"E{n:04d}",
        "operator_code": f"OP{n:04d}",
        "roles": ["cashier"],
        "privileges": ["checkin"],
    }
    base.update(kw)
    return base


def test_plugin_listed_menu_perms(opera_app):
    client, maker, _ = opera_app
    h = login(client)
    r = client.get("/api/v1/core/plugins", headers=h)
    assert r.status_code == 200
    opera = [p for p in r.json() if p["name"] == "opera_cloud"]
    assert opera and opera[0]["version"] == "1.0.0"
    assert set(opera[0]["permissions"]) == OPERA_PERMS
    r = client.get("/api/v1/core/menu", headers=h)
    assert r.status_code == 200
    assert any(i["path"] == "/opera" for i in r.json())
    db = maker()
    codes = {p.code for p in db.query(core_models.Permission).all()}
    db.close()
    assert OPERA_PERMS <= codes


def test_migration_table_exists(opera_app):
    _, _, test_engine = opera_app
    with test_engine.connect() as c:
        names = {row[0] for row in c.execute(text(
            "SELECT tablename FROM pg_tables WHERE schemaname = current_schema()"))}
        ver = c.execute(text("SELECT version_num FROM alembic_version_opera_cloud")).fetchall()
    assert "opera_accounts" in names
    assert ver and ver[0][0] == "opera_0001"


def test_rbac(opera_app):
    client, maker, _ = opera_app
    assert client.get(API).status_code == 401
    db = maker()
    db.add(core_models.User(username="noview", email="nv@x.io",
                            password_hash=auth.hash_password("Noview123!")))
    db.commit()
    db.close()
    vh = login(client, "noview", "Noview123!")
    assert client.get(API, headers=vh).status_code == 403
    assert client.get(API, headers=login(client)).status_code == 200


def test_crud_search_filters_audit(opera_app):
    client, maker, _ = opera_app
    h = login(client)
    # seed dept/position for FK checks
    d = client.post("/api/v1/core/departments", headers=h, json={"name": "FO"}).json()
    p = client.post("/api/v1/core/positions", headers=h,
                    json={"name": "Recep", "department_id": d["id"]}).json()
    r = client.post(API, headers=h, json=_payload(11, department_id=d["id"], position_id=p["id"]))
    assert r.status_code == 201, r.text
    acc = r.json()
    assert acc["department_name"] == "FO" and acc["position_name"] == "Recep"
    aid = acc["id"]
    client.post(API, headers=h, json=_payload(12))
    # q search across all five fields
    for q in ("op11", "op11@x.io", "Emp 11", "E0011", "OP0011"):
        r = client.get(API, headers=h, params={"q": q})
        assert r.status_code == 200 and r.json()["total"] >= 1, q
    # filters
    r = client.get(API, headers=h, params={"department_id": d["id"]})
    assert r.status_code == 200 and all(
        a["department_id"] == d["id"] for a in r.json()["items"])
    r = client.get(API, headers=h, params={"account_status": "active"})
    assert r.status_code == 200 and r.json()["total"] >= 2
    # pagination shape
    r = client.get(API, headers=h, params={"page": 2, "page_size": 1})
    body = r.json()
    assert set(body) == {"items", "total", "page", "page_size"}
    # get one + update + audit trail
    assert client.get(f"{API}/{aid}", headers=h).status_code == 200
    r = client.put(f"{API}/{aid}", headers=h,
                   json={"employee_name": "Emp Eleven", "account_status": "disabled"})
    assert r.status_code == 200 and r.json()["employee_name"] == "Emp Eleven"
    r = client.get(API, headers=h, params={"account_status": "disabled"})
    assert r.json()["total"] >= 1
    # delete
    assert client.delete(f"{API}/{aid}", headers=h).status_code == 204
    assert client.get(f"{API}/{aid}", headers=h).status_code == 404
    db = maker()
    rows = db.query(core_models.AuditLog).filter(
        core_models.AuditLog.resource == "opera_account").all()
    db.close()
    actions = {x.action for x in rows}
    assert {"create", "update", "delete"} <= actions


def test_validation_conflicts(opera_app):
    client, _, _ = opera_app
    h = login(client)
    client.post(API, headers=h, json=_payload(21))
    # duplicate username → 409
    r = client.post(API, headers=h, json=_payload(22, username="op21"))
    assert r.status_code == 409, r.text
    # duplicate operator_code → 409
    r = client.post(API, headers=h, json=_payload(23, operator_code="OP0021"))
    assert r.status_code == 409
    # bad email → 422
    r = client.post(API, headers=h, json=_payload(24, email="not-an-email"))
    assert r.status_code == 422
    # bad FK → 404
    r = client.post(API, headers=h, json=_payload(25, department_id=999999))
    assert r.status_code == 404
    # missing required → 422
    bad = _payload(26)
    del bad["operator_code"]
    assert client.post(API, headers=h, json=bad).status_code == 422


def test_export_xlsx(opera_app):
    client, _, _ = opera_app
    h = login(client)
    client.post(API, headers=h, json=_payload(31))
    r = client.get(f"{API}/export", headers=h)
    assert r.status_code == 200, r.text
    assert "spreadsheetml.sheet" in r.headers["content-type"]
    assert r.content[:2] == b"PK"  # xlsx is a zip
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert header == opera_services.EXPORT_COLUMNS
    assert ws.max_row >= 2
    # filter respected: export with q matching nothing has header only
    r = client.get(f"{API}/export", headers=h, params={"q": "zzz-no-such-user"})
    wb2 = load_workbook(io.BytesIO(r.content))
    assert wb2.active.max_row == 1


def test_expire_job(opera_app):
    client, maker, _ = opera_app
    h = login(client)
    yesterday = (date.today() - timedelta(days=1)).isoformat()
    r = client.post(API, headers=h, json=_payload(41, expiration_date=yesterday))
    assert r.status_code == 201, r.text
    aid = r.json()["id"]
    db = maker()
    n = opera_services.expire_old_accounts(db)
    db.close()
    assert n >= 1
    r = client.get(f"{API}/{aid}", headers=h)
    assert r.json()["account_status"] == "expired"